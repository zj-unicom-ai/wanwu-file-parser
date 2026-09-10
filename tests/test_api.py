"""API integration tests via FastAPI TestClient."""
from __future__ import annotations

import base64
import os
import io
from unittest.mock import patch

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.main import create_app
from app.services import parser


@pytest.fixture()
def client():
    return TestClient(create_app())


@pytest.fixture(autouse=True)
def _reset_client():
    parser.reset_client()
    yield
    parser.reset_client()


class TestHealth:
    def test_health(self, client):
        resp = client.get("/rag/health")
        assert resp.status_code == 200
        body = resp.json()
        assert body["status"] == "healthy"


class TestValidation:
    def test_missing_file(self, client):
        resp = client.post("/rag/model_parser_file")
        assert resp.status_code == 422  # FastAPI form validation (file required)

    def test_bad_extension(self, client):
        resp = client.post(
            "/rag/model_parser_file",
            files={"file": ("a.txt", b"x", "text/plain")},
            data={"file_name": "a.txt"},
        )
        assert resp.status_code == 400
        assert resp.json()["status"] == "failed"

    def test_path_traversal(self, client):
        resp = client.post(
            "/rag/model_parser_file",
            files={"file": ("../evil.pdf", b"x", "application/pdf")},
            data={"file_name": "../evil.pdf"},
        )
        assert resp.status_code == 400
        assert "非法字符" in resp.json()["message"]


class TestExcelShortcut:
    def _make_xlsx(self, path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["姓名", "年龄"])
        ws.append(["张三", "30"])
        wb.save(path)
        return path

    def test_xlsx_no_image_short_circuits(self, client, tmp_path):
        # The Excel shortcut must NOT call the OCR model.
        p = self._make_xlsx(tmp_path / "plain.xlsx")
        with patch("app.services.parser.get_client") as mock_get:
            with open(p, "rb") as fh:
                resp = client.post(
                    "/rag/model_parser_file",
                    files={"file": ("plain.xlsx", fh.read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                    data={"file_name": "plain.xlsx"},
                )
            # Model client must never be instantiated/used for no-image xlsx.
            mock_get.assert_not_called()
        assert resp.status_code == 200
        body = resp.json()
        assert body["status"] == "success"
        assert "张三" in body["content"]
        # prefix_image_url matches the config default (may be empty if unset).

    def test_xlsx_with_image_ocrs_image_and_keeps_text(self, client, tmp_path):
        import base64

        from openpyxl.drawing.image import Image as XLImage

        png = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNkYAAAAAYAAjCB0C8AAAAASUVORK5CYII="
        )
        p = tmp_path / "mixed.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["姓名", "年龄"])
        ws.append(["张三", "30"])
        ws.add_image(XLImage(io.BytesIO(png)), "A5")
        wb.save(p)

        # Stub PaddleOCRVLClient so no real OCR service is needed. The client is
        # imported lazily inside ocr_excel_images, so patch at its definition site.
        with patch("app.models.paddleocrvl.client.PaddleOCRVLClient") as MockClient:
            MockClient.return_value.parse_file.return_value = {
                "code": 200,
                "message": "ok",
                "data": {"md_content": "<IMG-OCR>", "json_data": "", "images": {}},
            }
            with open(p, "rb") as fh:
                resp = client.post(
                    "/rag/model_parser_file",
                    files={"file": ("mixed.xlsx", fh.read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                    data={"file_name": "mixed.xlsx"},
                )
        assert resp.status_code == 200
        body = resp.json()
        assert "张三" in body["content"]      # text kept
        assert "<IMG-OCR>" in body["content"]  # image OCR appended
        # The image OCR path was exercised exactly once (one image).
        assert MockClient.return_value.parse_file.call_count == 1


class TestExcelImageOcrEdgeCases:
    """Edge cases for the xlsx-with-images branch of ocr_excel_images.

    Covers the fault-tolerance paths the happy-path test does not: many images,
    a single image's OCR failing (others must still contribute), OCR returning
    empty markdown, and the client itself failing to load.
    """

    _PNG = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNkYAAAAAYAAjCB0C8AAAAASUVORK5CYII="
    )

    def _make_xlsx_with_n_images(self, path, n):
        """Build an xlsx with one text row + n embedded 1x1 PNGs."""
        from openpyxl.drawing.image import Image as XLImage

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["姓名", "年龄"])
        ws.append(["张三", "30"])
        for i in range(n):
            ws.add_image(XLImage(io.BytesIO(self._PNG)), f"A{5 + i}")
        wb.save(path)
        return path

    def _post(self, client, path):
        with open(path, "rb") as fh:
            return client.post(
                "/rag/model_parser_file",
                files={
                    "file": (
                        "mixed.xlsx",
                        fh.read(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
                data={"file_name": "mixed.xlsx"},
            )

    def test_multiple_images_all_ocrd_and_concatenated(self, client, tmp_path):
        # Three images -> three OCR calls, results joined by blank line.
        p = self._make_xlsx_with_n_images(tmp_path / "multi.xlsx", 3)
        with patch("app.models.paddleocrvl.client.PaddleOCRVLClient") as MockClient:
            MockClient.return_value.parse_file.side_effect = [
                {"code": 200, "message": "ok",
                 "data": {"md_content": "IMG1", "json_data": "", "images": {}}},
                {"code": 200, "message": "ok",
                 "data": {"md_content": "IMG2", "json_data": "", "images": {}}},
                {"code": 200, "message": "ok",
                 "data": {"md_content": "IMG3", "json_data": "", "images": {}}},
            ]
            resp = self._post(client, p)
        assert resp.status_code == 200
        body = resp.json()
        assert "张三" in body["content"]
        assert "IMG1" in body["content"]
        assert "IMG2" in body["content"]
        assert "IMG3" in body["content"]
        assert MockClient.return_value.parse_file.call_count == 3

    def test_one_image_ocr_failure_others_still_used(self, client, tmp_path):
        # Middle image raises; first and third must still contribute their text.
        p = self._make_xlsx_with_n_images(tmp_path / "partial.xlsx", 3)
        with patch("app.models.paddleocrvl.client.PaddleOCRVLClient") as MockClient:
            MockClient.return_value.parse_file.side_effect = [
                {"code": 200, "message": "ok",
                 "data": {"md_content": "GOOD1", "json_data": "", "images": {}}},
                RuntimeError("ocr service down for one image"),
                {"code": 200, "message": "ok",
                 "data": {"md_content": "GOOD3", "json_data": "", "images": {}}},
            ]
            resp = self._post(client, p)
        assert resp.status_code == 200
        body = resp.json()
        # Text preserved + the two successful OCRs concatenated; failed one skipped.
        assert "张三" in body["content"]
        assert "GOOD1" in body["content"]
        assert "GOOD3" in body["content"]
        assert MockClient.return_value.parse_file.call_count == 3

    def test_image_ocr_returns_empty_md_keeps_text_only(self, client, tmp_path):
        # OCR succeeds (code 200) but yields empty md -> warning, text still returned.
        p = self._make_xlsx_with_n_images(tmp_path / "empty.xlsx", 1)
        with patch("app.models.paddleocrvl.client.PaddleOCRVLClient") as MockClient:
            MockClient.return_value.parse_file.return_value = {
                "code": 200, "message": "ok",
                "data": {"md_content": "   ", "json_data": "", "images": {}},
            }
            resp = self._post(client, p)
        assert resp.status_code == 200
        body = resp.json()
        assert "张三" in body["content"]
        assert MockClient.return_value.parse_file.call_count == 1

    def test_image_ocr_non_200_code_skipped_keeps_text(self, client, tmp_path):
        # OCR returns code != 200 (e.g. 500) -> that image contributes nothing,
        # cell text is still returned.
        p = self._make_xlsx_with_n_images(tmp_path / "err500.xlsx", 1)
        with patch("app.models.paddleocrvl.client.PaddleOCRVLClient") as MockClient:
            MockClient.return_value.parse_file.return_value = {
                "code": 500, "message": "model error", "data": {},
            }
            resp = self._post(client, p)
        assert resp.status_code == 200
        body = resp.json()
        assert "张三" in body["content"]
        assert MockClient.return_value.parse_file.call_count == 1

    def test_paddleocrvl_client_import_failure_keeps_text(self, client, tmp_path):
        # If PaddleOCRVLClient cannot be imported (missing dep), ocr_excel_images
        # returns "" and the cell text is still returned (text-only degradation).
        p = self._make_xlsx_with_n_images(tmp_path / "noimport.xlsx", 1)
        with patch(
            "app.models.paddleocrvl.client.PaddleOCRVLClient",
            side_effect=ImportError("paddleocrvl deps missing"),
        ):
            resp = self._post(client, p)
        assert resp.status_code == 200
        body = resp.json()
        assert "张三" in body["content"]


class TestLegacyXls:
    """Best-effort .xls smoke test (legacy Excel 2003 format, not guaranteed)."""

    def test_xls_returns_text_without_pdf_or_ocr(self, client, tmp_path):
        # .xls (best-effort) is parsed to a markdown table in-CPU: no Stirling/PDF
        # and no OCR model call. Users should convert to .xlsx for full support.
        import xlwt

        p = tmp_path / "legacy.xls"
        wb = xlwt.Workbook()
        ws = wb.add_sheet("销售")
        ws.write(0, 0, "产品")
        ws.write(0, 1, "数量")
        ws.write(1, 0, "苹果")
        ws.write(1, 1, "12")
        wb.save(p)

        with patch("app.services.parser.get_client") as mock_get:
            with open(p, "rb") as fh:
                resp = client.post(
                    "/rag/model_parser_file",
                    files={"file": ("legacy.xls", fh.read(), "application/vnd.ms-excel")},
                    data={"file_name": "legacy.xls"},
                )
            # No OCR model call for .xls (text-only shortcut).
            mock_get.assert_not_called()
        assert resp.status_code == 200
        body = resp.json()
        assert body["status"] == "success"
        assert "苹果" in body["content"]
        assert "## 销售" in body["content"]


class TestOfficeNoLibreOffice:
    """When LibreOffice is not installed, Office files get a clear 400 error."""

    def test_docx_rejected_without_libreoffice(self, client, monkeypatch):
        from app.config import settings
        from app.services import file_convert

        monkeypatch.setattr(settings, "model_type", "paddleocrvl")
        # Simulate: LibreOffice is not installed
        monkeypatch.setattr(file_convert, "_find_libreoffice", lambda: None)
        with patch("app.services.parser.get_client"):
            resp = client.post(
                "/rag/model_parser_file",
                files={"file": ("doc.docx", b"%PDF-phony", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                data={"file_name": "doc.docx"},
            )
        assert resp.status_code == 400
        body = resp.json()
        assert body["code"] == "400"
        assert body["status"] == "failed"
        assert "LibreOffice" in body["message"] or "mineru" in body["message"]


class TestOfficeDoc2mdModes:
    """Tests for OFFICE_PROCESSING_MODE switching (auto / direct_extract / convert_pdf)."""

    def test_auto_mode_falls_back_to_convert_pdf_when_doc2md_fails(
        self, client, monkeypatch, tmp_path
    ):
        """In auto mode, if doc2md fails, fall back to LibreOffice->PDF->OCR."""
        from app.config import settings
        from app.services import file_convert

        monkeypatch.setattr(settings, "model_type", "paddleocrvl")
        monkeypatch.setattr(settings, "office_processing_mode", "auto")
        # doc2md CLI not installed -> doc2md fails
        monkeypatch.setattr(file_convert, "_find_paddleocr_cli", lambda: None)
        # LibreOffice is installed
        monkeypatch.setattr(file_convert, "_find_libreoffice", lambda: "/usr/bin/libreoffice")

        docx_path = tmp_path / "test.docx"
        docx_path.write_bytes(b"fake docx")

        def fake_lo_convert(file_path):
            pdf_path = file_path.replace(".docx", ".pdf")
            with open(pdf_path, "w") as f:
                f.write("%PDF-1.4 fake")
            return pdf_path

        monkeypatch.setattr(file_convert, "_try_libreoffice_convert", fake_lo_convert)

        with patch("app.services.parser.get_client") as mock_get:
            mock_get.return_value.parse_file.return_value = {
                "code": 200, "message": "ok",
                "data": {"md_content": "OCR content", "json_data": "", "images": {}},
            }
            mock_get.return_value.post_process.return_value = (
                "OCR content", "", ""
            )
            with open(docx_path, "rb") as fh:
                resp = client.post(
                    "/rag/model_parser_file",
                    files={"file": ("test.docx", fh.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                    data={"file_name": "test.docx"},
                )
        assert resp.status_code == 200
        assert "OCR content" in resp.json()["content"]

    def test_direct_extract_mode_raises_when_doc2md_not_installed(
        self, client, monkeypatch, tmp_path
    ):
        """In direct_extract mode, if paddleocr CLI is missing, return 400."""
        from app.config import settings
        from app.services import file_convert

        monkeypatch.setattr(settings, "model_type", "paddleocrvl")
        monkeypatch.setattr(settings, "office_processing_mode", "direct_extract")
        monkeypatch.setattr(file_convert, "_find_paddleocr_cli", lambda: None)

        docx_path = tmp_path / "test.docx"
        docx_path.write_bytes(b"fake docx")

        with patch("app.services.parser.get_client"):
            with open(docx_path, "rb") as fh:
                resp = client.post(
                    "/rag/model_parser_file",
                    files={"file": ("test.docx", fh.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                    data={"file_name": "test.docx"},
                )
        assert resp.status_code == 400
        body = resp.json()
        assert body["status"] == "failed"
        assert "paddleocr" in body["message"] or "doc2md" in body["message"]

    def test_convert_pdf_mode_skips_doc2md(self, client, monkeypatch, tmp_path):
        """In convert_pdf mode, doc2md is never tried; LibreOffice is used directly."""
        from app.config import settings
        from app.services import file_convert

        monkeypatch.setattr(settings, "model_type", "paddleocrvl")
        monkeypatch.setattr(settings, "office_processing_mode", "convert_pdf")
        monkeypatch.setattr(file_convert, "_find_paddleocr_cli", lambda: "/usr/bin/paddleocr")
        monkeypatch.setattr(file_convert, "_find_libreoffice", lambda: "/usr/bin/libreoffice")

        docx_path = tmp_path / "test.docx"
        docx_path.write_bytes(b"fake docx")

        def fake_lo_convert(file_path):
            pdf_path = file_path.replace(".docx", ".pdf")
            with open(pdf_path, "w") as f:
                f.write("%PDF-1.4 fake")
            return pdf_path

        monkeypatch.setattr(file_convert, "_try_libreoffice_convert", fake_lo_convert)

        with patch("app.services.parser.get_client") as mock_get:
            mock_get.return_value.parse_file.return_value = {
                "code": 200, "message": "ok",
                "data": {"md_content": "PDF OCR content", "json_data": "", "images": {}},
            }
            mock_get.return_value.post_process.return_value = (
                "PDF OCR content", "", ""
            )
            with open(docx_path, "rb") as fh:
                resp = client.post(
                    "/rag/model_parser_file",
                    files={"file": ("test.docx", fh.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                    data={"file_name": "test.docx"},
                )
        assert resp.status_code == 200
        assert "PDF OCR content" in resp.json()["content"]

    def test_auto_mode_uses_doc2md_when_available(self, client, monkeypatch, tmp_path):
        """In auto mode, when paddleocr CLI is installed, doc2md is used."""
        from app.config import settings
        from app.services import file_convert

        monkeypatch.setattr(settings, "model_type", "paddleocrvl")
        monkeypatch.setattr(settings, "office_processing_mode", "auto")
        monkeypatch.setattr(file_convert, "_find_paddleocr_cli", lambda: "/usr/bin/paddleocr")

        docx_path = tmp_path / "test.docx"
        docx_path.write_bytes(b"fake docx")

        # Mock the doc2md subprocess to produce a .md file
        def fake_doc2md_run(cmd, **kwargs):
            out_dir = cmd[cmd.index("-o") + 1]
            base = os.path.splitext(os.path.basename(cmd[cmd.index("-i") + 1]))[0]
            md_path = os.path.join(out_dir, f"{base}.md")
            with open(md_path, "w") as f:
                f.write("# Title\n\ndoc2md extracted content")

            class _Result:
                returncode = 0
                stdout = ""
                stderr = ""

            return _Result()

        with patch("subprocess.run", side_effect=fake_doc2md_run):
            with patch("app.services.parser.get_client") as mock_get:
                with open(docx_path, "rb") as fh:
                    resp = client.post(
                        "/rag/model_parser_file",
                        files={"file": ("test.docx", fh.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                        data={"file_name": "test.docx"},
                    )
        assert resp.status_code == 200
        assert "doc2md extracted content" in resp.json()["content"]
        # Model client was NOT called (doc2md shortcut returned directly)
        mock_get.assert_not_called()
