"""Excel text + image extraction tests (the .xlsx text+image split feature)."""
from __future__ import annotations

import base64
import io
import os

import openpyxl
from openpyxl.drawing.image import Image as XLImage

from app.services.excel_service import (
    excel_to_markdown,
    extract_embedded_images,
)

# 1x1 transparent PNG.
_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNkYAAAAAYAAjCB0C8AAAAASUVORK5CYII="
)


class TestExtractEmbeddedImages:
    def test_xls_returns_none_sentinel(self, tmp_path):
        # .xls branch does not read the file; returns None sentinel.
        assert extract_embedded_images(str(tmp_path / "no-such-file.xls")) is None

    def test_xlsx_without_images_returns_empty_list(self, tmp_path):
        p = tmp_path / "plain.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["a", "b"])
        wb.save(p)
        assert extract_embedded_images(str(p)) == []

    def test_xlsx_with_image_extracts_temp_file(self, tmp_path):
        p = tmp_path / "with_img.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["a", "b"])
        ws.add_image(XLImage(io.BytesIO(_PNG)), "A5")
        wb.save(p)

        images = extract_embedded_images(str(p))
        assert len(images) == 1
        assert os.path.exists(images[0])
        assert os.path.getsize(images[0]) > 0
        assert images[0].endswith(".png")
        os.remove(images[0])


class TestExcelToMarkdown:
    def test_xls_returns_none_images(self, tmp_path):
        md, has_images, image_paths = excel_to_markdown(str(tmp_path / "any.xls"))
        assert has_images is True
        assert md == ""
        assert image_paths is None

    def test_xlsx_no_image_short_circuits(self, tmp_path):
        p = tmp_path / "plain.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["姓名", "年龄"])
        ws.append(["张三", "30"])
        wb.save(p)

        md, has_images, image_paths = excel_to_markdown(str(p))
        assert has_images is False
        assert image_paths == []
        assert "张三" in md
        assert "30" in md

    def test_xlsx_with_image_keeps_text_and_extracts_image(self, tmp_path):
        p = tmp_path / "mixed.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["姓名", "年龄"])
        ws.append(["张三", "30"])
        ws.add_image(XLImage(io.BytesIO(_PNG)), "A5")
        wb.save(p)

        md, has_images, image_paths = excel_to_markdown(str(p))
        assert has_images is True
        assert len(image_paths) == 1
        assert "张三" in md  # text preserved
        os.remove(image_paths[0])

    def test_pipe_in_cell_is_escaped(self, tmp_path):
        p = tmp_path / "pipes.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["col1"])      # header row
        ws.append(["a|b"])        # data row with a pipe
        wb.save(p)
        md, _, _ = excel_to_markdown(str(p))
        assert "a\\|b" in md
